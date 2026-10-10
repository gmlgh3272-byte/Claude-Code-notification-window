"""도면 분석 결과(JSON)와 내역서(JSON)를 품목매핑표 기준으로 대조.

사용법:
    python compare.py --drawing 도면1.json [도면2.json ...] --boq 내역.json
                      --map 품목매핑.csv --out 대조결과.xlsx [--len-range 1.0,1.5]
"""
import argparse
import csv
import fnmatch
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

TO_METER = {"mm": 0.001, "cm": 0.01, "m": 1.0, "km": 1000.0}


def norm(s):
    return re.sub(r"\s+", "", str(s or "")).upper()


def split(s):
    return [p.strip() for p in str(s or "").split("|") if p.strip()]


def load_map(path):
    with open(path, encoding="utf-8-sig") as f:
        rows = []
        for r in csv.DictReader(f):
            if not r.get("품목"):
                continue
            rows.append({"item": r["품목"].strip(), "kind": (r.get("구분") or "개수").strip(),
                         "dwg": [p.upper() for p in split(r.get("도면패턴"))],
                         "boq_name": [norm(k) for k in split(r.get("내역품명"))],
                         "boq_spec": [norm(k) for k in split(r.get("내역규격"))],
                         "note": r.get("비고", "")})
        return rows


def dwg_match(rule, name):
    return any(fnmatch.fnmatch(name.upper(), p) for p in rule["dwg"])


def boq_match(rule, it):
    if not rule["boq_name"] or not any(k in norm(it["name"]) for k in rule["boq_name"]):
        return False
    return not rule["boq_spec"] or any(k in norm(it["spec"]) for k in rule["boq_spec"])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--drawing", nargs="+", required=True)
    ap.add_argument("--boq", required=True)
    ap.add_argument("--map", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--len-range", default="1.0,1.5", help="길이 항목 내역/도면 비율 허용범위 (할증·수직길이 감안)")
    args = ap.parse_args()
    lo, hi = (float(x) for x in args.len_range.split(","))

    rules = load_map(args.map)
    res = {r["item"]: {"rule": r, "dwg": 0.0, "boq": 0.0, "dwg_src": [], "boq_src": [], "units": set()} for r in rules}

    unmapped_blocks, drawings = [], []
    for path in args.drawing:
        d = json.loads(Path(path).read_text(encoding="utf-8"))
        drawings.append(d.get("file", Path(path).name))
        factor = TO_METER.get(d.get("units"), None)
        for b in d.get("blocks", []):
            rule = next((r for r in rules if r["kind"] == "개수" and dwg_match(r, b["block"])), None)
            if rule:
                res[rule["item"]]["dwg"] += b["count"]
                res[rule["item"]]["dwg_src"].append(f'{b["block"]}({b["count"]})')
            else:
                unmapped_blocks.append({"file": d.get("file"), "block": b["block"], "count": b["count"], "layers": b["layers"]})
        for layer in d.get("layers", []):
            if not layer.get("length"):
                continue
            rule = next((r for r in rules if r["kind"] == "길이" and dwg_match(r, layer["layer"])), None)
            if rule:
                meters = layer["length"] * (factor if factor else 0.001)
                res[rule["item"]]["dwg"] += meters
                res[rule["item"]]["dwg_src"].append(f'{layer["layer"]}({meters:.1f}m{"" if factor else ", 단위추정mm"})')

    boq = json.loads(Path(args.boq).read_text(encoding="utf-8"))
    unmapped_boq = []
    for it in boq["items"]:
        rule = next((r for r in rules if boq_match(r, it)), None)
        if rule:
            x = res[rule["item"]]
            x["boq"] += it["qty"]
            x["units"].add(it["unit"])
            x["boq_src"].append(f'{it["sheet"]}!{it["row"]} {it["name"]} {it["spec"]}({it["qty"]:g})')
        else:
            unmapped_boq.append(it)

    rows = []
    for item, x in res.items():
        kind, dq, bq = x["rule"]["kind"], round(x["dwg"], 2), round(x["boq"], 2)
        if dq == 0 and bq == 0:
            verdict = "해당없음"
        elif kind == "개수":
            verdict = "일치" if dq == bq else ("도면에만 있음" if bq == 0 else "내역에만 있음" if dq == 0 else "수량 불일치")
        else:
            ratio = bq / dq if dq else None
            verdict = "도면 길이 없음(확인)" if ratio is None else ("범위 내(참고)" if lo <= ratio <= hi else "범위 밖(확인)")
        rows.append({"품목": item, "구분": kind, "도면수량": dq, "내역수량": bq, "차이(내역-도면)": round(bq - dq, 2),
                     "판정": verdict, "내역단위": ", ".join(sorted(u for u in x["units"] if u)),
                     "도면근거": "; ".join(x["dwg_src"]), "내역근거": "; ".join(x["boq_src"][:30]), "비고": x["rule"]["note"]})

    wb = Workbook()
    ws = wb.active
    ws.title = "대조표"
    header = list(rows[0].keys()) if rows else ["품목"]
    ws.append(header)
    red = PatternFill("solid", fgColor="FFC7CE")
    yellow = PatternFill("solid", fgColor="FFEB9C")
    for r in rows:
        ws.append([r[h] for h in header])
        fill = red if r["판정"] in ("수량 불일치", "도면에만 있음", "내역에만 있음", "범위 밖(확인)", "도면 길이 없음(확인)") else \
            yellow if r["판정"] == "범위 내(참고)" else None
        if fill:
            for c in ws[ws.max_row]:
                c.fill = fill

    def sheet(title, hdr, data):
        s = wb.create_sheet(title)
        s.append(hdr)
        for d in data:
            s.append([d.get(h) for h in hdr])

    sheet("도면_미매핑블록", ["file", "block", "count", "layers"], unmapped_blocks)
    sheet("내역_미매핑품목", ["sheet", "row", "section", "name", "spec", "unit", "qty"], unmapped_boq)
    sheet("설정", ["항목", "값"], [{"항목": "도면", "값": ", ".join(drawings)}, {"항목": "내역서", "값": ", ".join(boq["files"])},
                                {"항목": "길이 허용비율", "값": args.len_range}])
    for s in wb.worksheets:
        for c in s[1]:
            c.font = Font(bold=True)
        s.freeze_panes = "A2"
        for col in s.columns:
            w = max(len(str(c.value or "")) for c in col[:200])
            s.column_dimensions[col[0].column_letter].width = min(max(w * 1.2, 8), 50)
    wb.save(args.out)

    print(json.dumps({"out": str(Path(args.out).resolve()),
                      "rows": [{k: r[k] for k in ["품목", "구분", "도면수량", "내역수량", "판정"]} for r in rows],
                      "unmapped_blocks": sorted(unmapped_blocks, key=lambda b: -b["count"])[:30],
                      "unmapped_boq_count": len(unmapped_boq),
                      "unmapped_boq_sample": [f'{i["name"]} {i["spec"]} {i["qty"]:g}{i["unit"]}' for i in unmapped_boq[:40]]},
                     ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()

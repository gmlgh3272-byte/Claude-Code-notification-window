"""엑셀 내역서(.xlsx) 읽기: 시트별 머리행 자동 탐지 후 품명/규격/단위/수량 추출.

사용법:
    python read_boq.py <내역서.xlsx 또는 폴더> --out 내역.json [--sheets 시트1,시트2]
"""
import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

NAME_KEYS = ["품명및규격", "공종및규격", "품명", "공종명", "명칭", "품목", "공종"]
SKIP_SHEET = re.compile(r"일위대가|단가산출|산출근거|호표|원가계산|집계|표지|갑지|목차")
TOTAL_ROW = re.compile(r"^[\[\(<]?(소계|합계|총계|누계|계|공급가액|부가가치세|총공사비)[\]\)>]?$")


def norm(v):
    return re.sub(r"\s+", "", str(v)) if v is not None else ""


def to_num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def find_header(ws, scan=40):
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=scan, values_only=True), start=1):
        cells = [norm(c) for c in row]
        if "수량" not in cells:
            continue
        cols = {"qty": cells.index("수량")}
        for key in NAME_KEYS:
            if key in cells:
                cols["name"] = cells.index(key)
                cols["combined"] = key.endswith("규격")
                break
        if "name" not in cols:
            continue
        if "규격" in cells:
            cols["spec"] = cells.index("규격")
        if "단위" in cells:
            cols["unit"] = cells.index("단위")
        return r_idx, cols
    return None, None


def read_sheet(ws, file_name):
    header_row, cols = find_header(ws)
    if not header_row:
        return None
    items, section = [], ""
    for r_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        get = lambda k: row[cols[k]] if k in cols and cols[k] < len(row) else None
        name = str(get("name") or "").strip()
        if not name:
            continue
        qty = to_num(get("qty"))
        unit = str(get("unit") or "").strip()
        spec = str(get("spec") or "").strip()
        if TOTAL_ROW.match(norm(name)):
            continue
        if not qty and not unit:
            section = name  # 공종 제목 행
            continue
        if not qty:
            continue
        items.append({"file": file_name, "sheet": ws.title, "row": r_idx, "section": section,
                      "name": name, "spec": spec, "unit": unit, "qty": qty})
    return items


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("--out", required=True)
    ap.add_argument("--sheets", help="읽을 시트 이름을 쉼표로 지정 (자동 제외 규칙 무시)")
    args = ap.parse_args()

    src = Path(args.src)
    files = sorted(p for p in src.glob("*.xlsx") if not p.name.startswith("~$")) if src.is_dir() else [src]
    if src.suffix.lower() == ".xls" or (src.is_dir() and list(src.glob("*.xls"))):
        print("[안내] .xls(구버전)는 지원 안 함. 엑셀에서 .xlsx로 다른 이름 저장 필요")
    forced = set(args.sheets.split(",")) if args.sheets else None

    all_items, used, skipped = [], [], []
    for f in files:
        wb = load_workbook(f, data_only=True, read_only=True)
        for ws in wb.worksheets:
            if forced is not None and ws.title not in forced:
                continue
            if forced is None and SKIP_SHEET.search(ws.title):
                skipped.append(f"{f.name}:{ws.title} (이름 규칙으로 제외)")
                continue
            items = read_sheet(ws, f.name)
            if items is None:
                skipped.append(f"{f.name}:{ws.title} (머리행 없음)")
            else:
                used.append(f"{f.name}:{ws.title} ({len(items)}행)")
                all_items.extend(items)

    agg = defaultdict(float)
    for it in all_items:
        agg[(it["name"], it["spec"], it["unit"])] += it["qty"]
    grouped = [{"name": n, "spec": s, "unit": u, "qty": round(q, 3)} for (n, s, u), q in agg.items()]

    no_value = sum(1 for it in all_items if it["qty"] is None)
    result = {"files": [f.name for f in files], "sheets_used": used, "sheets_skipped": skipped,
              "item_rows": len(all_items), "grouped_items": len(grouped),
              "items": all_items, "grouped": grouped}
    Path(args.out).write_text(json.dumps(result, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({k: result[k] for k in ["files", "sheets_used", "sheets_skipped", "item_rows", "grouped_items"]}
                     | {"sample": grouped[:30], "out": str(Path(args.out).resolve())}, ensure_ascii=False, indent=1))
    if not all_items:
        print("[경고] 추출된 품목 없음. 수식 결과가 비어 있으면 엑셀에서 한 번 열고 저장 후 재시도")


if __name__ == "__main__":
    main()
